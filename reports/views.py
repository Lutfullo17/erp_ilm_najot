import openpyxl
from django.http import HttpResponse
from django.views import View
from django.contrib.auth.mixins import LoginRequiredMixin
from users1.views import AdminRequiredMixin

from students.models import Student
from payments.models import PaymentTransaction
from attendance.models import AttendanceRecord

class ExportStudentsView(LoginRequiredMixin, AdminRequiredMixin, View):
    def get(self, request):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "O'quvchilar"
        
        headers = ['ID', 'Fishi', 'Ismi', 'Telefon', 'Manzil', 'Holati']
        ws.append(headers)
        
        students = Student.objects.all().order_by('last_name')
        for student in students:
            ws.append([
                student.id,
                student.last_name,
                student.first_name,
                student.phone,
                student.address,
                'Faol' if student.is_active else 'Nofaol'
            ])
            
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename=students.xlsx'
        wb.save(response)
        return response

class ExportPaymentsView(LoginRequiredMixin, AdminRequiredMixin, View):
    def get(self, request):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "To'lovlar"
        
        headers = ['Sana', "O'quvchi", 'Guruh', 'Miqdor', 'Turi', 'Eslatma']
        ws.append(headers)
        
        payments = PaymentTransaction.objects.all().select_related('student', 'group').order_by('-payment_date')
        for payment in payments:
            ws.append([
                payment.payment_date.strftime('%Y-%m-%d'),
                str(payment.student),
                payment.group.name,
                payment.amount,
                payment.get_method_display(),
                payment.note
            ])
            
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename=payments.xlsx'
        wb.save(response)
        return response

class ExportAttendanceView(LoginRequiredMixin, AdminRequiredMixin, View):
    def get(self, request):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Davomat"
        
        headers = ['Sana', 'Guruh', "O'quvchi", 'Holat', 'Izoh']
        ws.append(headers)
        
        records = AttendanceRecord.objects.all().select_related('session', 'session__group', 'student').order_by('-session__date')
        for record in records:
            ws.append([
                record.session.date.strftime('%Y-%m-%d'),
                record.session.group.name,
                str(record.student),
                record.get_status_display(),
                record.comment
            ])
            
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename=attendance.xlsx'
        wb.save(response)
        return response
